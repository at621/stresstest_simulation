"""The run log and the Excel report (R10, R11): what happened is visible and reproducible."""

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from stress.contracts.run import RUN_LOG
from stress.lineage import compare_outputs, sha256_file
from stress.report import WORKBOOK_FILE
from tests import builders
from tests.support import tolerance

SHEETS = ("Summary", "Portfolios", "Results by portfolio", "Impairment", "Stage migration",
          "Coverage by stage", "S3 cohorts", "Parameters", "Starting point", "Checks", "Notes")


@pytest.mark.req("OPS-01", "OPS-07")
def test_run_log_records_every_step(core_run):
    ctx, manifest = core_run
    log = (ctx.run_dir / RUN_LOG).read_text(encoding="utf-8")
    for step in ("R1", "R2", "R3", "R4", "R5", "R6", "R7", "R8", "R9", "R10", "R11"):
        assert f"stress.{step} " in log, f"{step} missing from the run log"
    assert "start: load inputs" in log and "done in" in log
    assert "control totals match the manifest" in log
    assert "no record rejected" in log
    assert "sealing the run as SYNTHETIC" in log
    # The log has times, so it is hashed for integrity but is not an output a rerun must match.
    assert manifest["log"] == {RUN_LOG: sha256_file(ctx.run_dir / RUN_LOG)}
    assert RUN_LOG not in "".join(manifest["outputs"])


@pytest.mark.req("OPS-07")
def test_blocked_run_log_names_the_reason(mini, tmp_path):
    mini["scenarios"]["scenarios"][1]["status"] = "OFFICIAL"
    ctx, manifest = builders.run(builders.write(mini, tmp_path / "in"), tmp_path / "runs")
    log = (ctx.run_dir / RUN_LOG).read_text(encoding="utf-8")
    assert manifest["status"] == "BLOCKED"
    assert "BLOCKED SCENARIO_PROVENANCE" in log and "sealing the run as BLOCKED" in log


@pytest.mark.req("RPT-05")
def test_excel_report_matches_the_engine(core_run):
    ctx, manifest = core_run
    path = ctx.reports_dir / WORKBOOK_FILE
    assert f"results/reports/{WORKBOOK_FILE}" in manifest["outputs"]
    wb = load_workbook(path)
    assert tuple(wb.sheetnames) == SHEETS
    assert wb.calculation.fullCalcOnLoad, "Excel recalculates every formula on open"
    # Engine values on the Impairment sheet equal impairment.csv, in EUR millions.
    imp = builders.report(ctx, "impairment.csv")
    ws = wb["Impairment"]
    header = [c.value for c in ws[4]]
    rows = pd.DataFrame([[c.value for c in r] for r in ws.iter_rows(min_row=5)], columns=header)
    merged = rows.merge(imp.assign(Year=imp["year"].astype(str)),
                        left_on=["Scenario", "Year", "Portfolio"], right_on=["scenario", "Year", "portfolio"])
    assert len(merged) == len(imp)
    tol = tolerance("amount")["abs"]
    assert np.allclose(merged["Closing allowance (on)"] * 1e6, merged["closing_on"], atol=tol, rtol=0)
    assert np.allclose(merged["Cumulative charge"] * 1e6, merged["cumulative_charge"], atol=tol, rtol=0)
    # Derived figures are formulas, never pasted results.
    assert str(ws.cell(row=5, column=header.index("Closing total") + 1).value).startswith("=")
    summary = wb["Summary"]
    formulas = [c.value for row in summary.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("SUMIFS(" in f for f in formulas) and len(formulas) > 20
    fonts = {c.font.name for row in summary.iter_rows() for c in row if c.value is not None}
    assert fonts == {"Arial"}


@pytest.mark.control("Reproducibility")
def test_excel_report_is_byte_reproducible(mini_tables, tmp_path):
    inputs = builders.write(builders.copy_tables(mini_tables), tmp_path / "in")
    first, a = builders.run(inputs, tmp_path / "runs")
    second, b = builders.run(inputs, tmp_path / "runs")
    assert first.run_id != second.run_id
    key = f"results/reports/{WORKBOOK_FILE}"
    assert a["outputs"][key] == b["outputs"][key]
    assert compare_outputs(a, b) == []
